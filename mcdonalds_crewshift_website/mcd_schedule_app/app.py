import os, csv, io, re, sqlite3, secrets, hashlib, hmac, json, calendar
from datetime import datetime, date, timedelta
from functools import wraps
from pathlib import Path
from flask import Flask, render_template, request, redirect, url_for, flash, session, jsonify, abort, send_file
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from urllib.request import Request, urlopen

BASE=Path(__file__).parent
app=Flask(__name__)
app.secret_key=os.getenv('SECRET_KEY') or 'dev-only-change-before-deployment-'+secrets.token_hex(12)
app.config['MAX_CONTENT_LENGTH']=8*1024*1024
DB=os.getenv('DATABASE_PATH',str(BASE/'schedule.db'))
SHEET_ID='1WFqS2dwMiESsFgGnLowfxCWYTJQCgREv521vuv_O4i0'


def conn():
    c=sqlite3.connect(DB); c.row_factory=sqlite3.Row; c.execute('PRAGMA foreign_keys=ON'); return c

def init_db():
    with conn() as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY, employee_id TEXT UNIQUE, full_name TEXT NOT NULL, username TEXT NOT NULL UNIQUE, password_hash TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('employee','manager')), active INTEGER NOT NULL DEFAULT 1, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS shifts(id INTEGER PRIMARY KEY, employee_id INTEGER NOT NULL REFERENCES users(id), work_date TEXT NOT NULL, start_time TEXT NOT NULL, end_time TEXT NOT NULL, break_minutes INTEGER NOT NULL DEFAULT 0, station TEXT DEFAULT '', category TEXT DEFAULT '', status TEXT DEFAULT 'Scheduled', notes TEXT DEFAULT '', source TEXT DEFAULT 'manual', source_key TEXT UNIQUE, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE INDEX IF NOT EXISTS idx_shift_date ON shifts(work_date);
        CREATE TABLE IF NOT EXISTS time_off(id INTEGER PRIMARY KEY, employee_id INTEGER NOT NULL REFERENCES users(id), start_date TEXT NOT NULL, end_date TEXT NOT NULL, reason TEXT NOT NULL, status TEXT NOT NULL DEFAULT 'pending', manager_comment TEXT DEFAULT '', reviewed_by INTEGER REFERENCES users(id), reviewed_at TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS swaps(id INTEGER PRIMARY KEY, requester_id INTEGER NOT NULL REFERENCES users(id), target_id INTEGER NOT NULL REFERENCES users(id), requester_shift_id INTEGER NOT NULL REFERENCES shifts(id), target_shift_id INTEGER NOT NULL REFERENCES shifts(id), status TEXT NOT NULL DEFAULT 'pending', manager_comment TEXT DEFAULT '', target_accepted INTEGER DEFAULT NULL, reviewed_by INTEGER REFERENCES users(id), reviewed_at TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        CREATE TABLE IF NOT EXISTS imports(id INTEGER PRIMARY KEY, filename TEXT, sheet_name TEXT, imported INTEGER, skipped INTEGER, notes TEXT, created_at TEXT DEFAULT CURRENT_TIMESTAMP);
        ''')
        if not c.execute('SELECT id FROM users WHERE role="manager"').fetchone():
            username=os.getenv('ADMIN_USERNAME','admin'); pwd=os.getenv('ADMIN_PASSWORD')
            if not pwd:
                pwd=secrets.token_urlsafe(12)
                print('\n'+'='*65+'\nFIRST-RUN MANAGER LOGIN\nUsername: '+username+'\nPassword: '+pwd+'\nSave it now; change it after login.\n'+'='*65+'\n',flush=True)
            c.execute('INSERT INTO users(full_name,username,password_hash,role) VALUES(?,?,?,?)',('Store Manager',username,generate_password_hash(pwd),'manager'))

@app.before_request
def csrf_guard():
    if 'csrf' not in session: session['csrf']=secrets.token_hex(24)
    if request.method in ('POST','PUT','PATCH','DELETE'):
        if not hmac.compare_digest(request.form.get('csrf_token') or request.headers.get('X-CSRF-Token',''),session['csrf']): abort(400,'Invalid CSRF token')

@app.context_processor
def globals_for_templates():
    return {'me': user(), 'csrf_token': session.get('csrf',''), 'today':date.today().isoformat()}

def user():
    uid=session.get('uid')
    if not uid:return None
    with conn() as c:return c.execute('SELECT id,employee_id,full_name,username,role,active FROM users WHERE id=? AND active=1',(uid,)).fetchone()

def login_required(f):
    @wraps(f)
    def inner(*args,**kwargs):
        if not user():return redirect(url_for('login'))
        return f(*args,**kwargs)
    return inner

def manager_required(f):
    @wraps(f)
    def inner(*args,**kwargs):
        u=user()
        if not u:return redirect(url_for('login'))
        if u['role']!='manager': abort(403)
        return f(*args,**kwargs)
    return inner

def parse_date(s): return date.fromisoformat(s)
def parse_time(s):return datetime.strptime(s,'%H:%M').time()
def duration(start,end,br):
    a=datetime.combine(date.today(),parse_time(start)); b=datetime.combine(date.today(),parse_time(end))
    if b<=a:b+=timedelta(days=1)
    total=(b-a).total_seconds()/60-int(br)
    if total<=0 or total>1440:raise ValueError('Break must be shorter than the shift and shifts cannot exceed 24 hours.')
    return round(total/60,2)

def shift_rows(where='',params=()):
    with conn() as c:
        return [dict(x) for x in c.execute('''SELECT s.*,u.full_name,u.employee_id as employee_code,u.active FROM shifts s JOIN users u ON u.id=s.employee_id WHERE 1=1 '''+where+' ORDER BY s.work_date,s.start_time,u.full_name',params).fetchall()]

def week_start(s):
    try:d=parse_date(s)
    except (TypeError,ValueError):d=date.today()
    return d-timedelta(days=d.weekday())

def check_shift(c,emp,d,start,end,exclude=None):
    parse_date(d); duration(start,end,0)
    # Overlap check includes overnight portions carried into next day.
    begin=datetime.combine(parse_date(d),parse_time(start)); finish=datetime.combine(parse_date(d),parse_time(end))
    if finish<=begin:finish+=timedelta(days=1)
    options=c.execute('SELECT * FROM shifts WHERE employee_id=? AND work_date BETWEEN ? AND ?', (emp,(parse_date(d)-timedelta(days=1)).isoformat(),(parse_date(d)+timedelta(days=1)).isoformat())).fetchall()
    for s in options:
        if exclude and s['id']==exclude:continue
        a=datetime.combine(parse_date(s['work_date']),parse_time(s['start_time'])); b=datetime.combine(parse_date(s['work_date']),parse_time(s['end_time']))
        if b<=a:b+=timedelta(days=1)
        if begin<b and a<finish:raise ValueError('This employee already has an overlapping shift.')

def validate_shift_form(f):
    emp=int(f['employee_id']); d=f['work_date']; st=f['start_time']; en=f['end_time']; br=int(f.get('break_minutes') or 0)
    if br<0 or br>300:raise ValueError('Break must be between 0 and 300 minutes.')
    duration(st,en,br)
    with conn() as c:
        if not c.execute('SELECT id FROM users WHERE id=? AND active=1 AND role="employee"',(emp,)).fetchone(): raise ValueError('Select an active employee.')
    return (emp,d,st,en,br,f.get('station','').strip()[:80],f.get('category','').strip()[:40],f.get('status','Scheduled').strip()[:40],f.get('notes','').strip()[:500])

@app.route('/')
def root():return redirect(url_for('dashboard') if user() else url_for('login'))
@app.route('/login',methods=['GET','POST'])
def login():
    if user():return redirect(url_for('dashboard'))
    if request.method=='POST':
        with conn() as c:u=c.execute('SELECT * FROM users WHERE username=? AND active=1',(request.form.get('username','').strip(),)).fetchone()
        if u and check_password_hash(u['password_hash'],request.form.get('password','')):
            session.clear();session['uid']=u['id'];session['csrf']=secrets.token_hex(24);return redirect(url_for('dashboard'))
        flash('Incorrect username or password.','error')
    return render_template('login.html')
@app.post('/logout')
def logout():session.clear();return redirect(url_for('login'))
@app.route('/account',methods=['GET','POST'])
@login_required
def account():
    if request.method=='POST':
        with conn() as c:
            u=c.execute('SELECT * FROM users WHERE id=?',(user()['id'],)).fetchone()
            if not check_password_hash(u['password_hash'],request.form.get('current_password','')):flash('Current password is incorrect.','error')
            elif len(request.form.get('password',''))<10:flash('New password must be at least 10 characters.','error')
            else:c.execute('UPDATE users SET password_hash=? WHERE id=?',(generate_password_hash(request.form['password']),u['id']));flash('Password updated.','success')
        return redirect(url_for('account'))
    return render_template('account.html')
@app.route('/dashboard')
@login_required
def dashboard():
    u=user(); start=week_start(request.args.get('week'));end=start+timedelta(days=6)
    where=' AND s.work_date BETWEEN ? AND ?';p=[start.isoformat(),end.isoformat()]
    if u['role']=='employee':where+=' AND s.employee_id=?';p.append(u['id'])
    shifts=shift_rows(where,p)
    for s in shifts:s['hours']=duration(s['start_time'],s['end_time'],s['break_minutes'])
    with conn() as c:
        pending=c.execute('SELECT count(*) FROM time_off WHERE status="pending"'+(' AND employee_id=?' if u['role']=='employee' else ''),(() if u['role']=='manager' else (u['id'],))).fetchone()[0]
        pending+=c.execute('SELECT count(*) FROM swaps WHERE status="pending"'+(' AND (requester_id=? OR target_id=?)' if u['role']=='employee' else ''),(() if u['role']=='manager' else (u['id'],u['id']))).fetchone()[0]
        count=c.execute('SELECT count(*) FROM users WHERE role="employee" AND active=1').fetchone()[0]
    return render_template('dashboard.html',shifts=shifts,start=start,end=end,prev=(start-timedelta(days=7)).isoformat(),next=(start+timedelta(days=7)).isoformat(),pending=pending,employee_count=count,total_hours=round(sum(s['hours'] for s in shifts),1))
@app.route('/schedule')
@login_required
def schedule():
    u=user();view=request.args.get('view','week');start=week_start(request.args.get('week'))
    if view=='month':
        x=request.args.get('month') or date.today().strftime('%Y-%m')
        try:y,m=map(int,x.split('-'));first=date(y,m,1)
        except (ValueError,TypeError):first=date.today().replace(day=1)
        start=first-timedelta(days=first.weekday());end=(first.replace(day=calendar.monthrange(first.year,first.month)[1])+timedelta(days=6));end=end-timedelta(days=end.weekday())
    else:end=start+timedelta(days=6)
    filters={'employee':request.args.get('employee',''),'station':request.args.get('station',''),'search':request.args.get('search',''),'shift':request.args.get('shift','')}
    w=' AND s.work_date BETWEEN ? AND ?';p=[start.isoformat(),end.isoformat()]
    if u['role']=='employee':w+=' AND s.employee_id=?';p.append(u['id'])
    elif filters['employee'].isdigit():w+=' AND s.employee_id=?';p.append(int(filters['employee']))
    if filters['station']:w+=' AND LOWER(s.station) LIKE ?';p.append('%'+filters['station'].lower()+'%')
    if filters['shift']:w+=' AND s.category=?';p.append(filters['shift'])
    if filters['search']:w+=' AND (LOWER(u.full_name) LIKE ? OR LOWER(s.notes) LIKE ? OR LOWER(s.station) LIKE ?)';p.extend(['%'+filters['search'].lower()+'%']*3)
    shifts=shift_rows(w,p)
    for s in shifts:s['hours']=duration(s['start_time'],s['end_time'],s['break_minutes'])
    with conn() as c:employees=c.execute('SELECT id,full_name FROM users WHERE role="employee" AND active=1 ORDER BY full_name').fetchall()
    days=[start+timedelta(days=i) for i in range((end-start).days+1)];grouped={d.isoformat():[s for s in shifts if s['work_date']==d.isoformat()] for d in days}
    return render_template('schedule.html',shifts=shifts,days=days,grouped=grouped,start=start,end=end,view=view,filters=filters,employees=employees,prev=(start-timedelta(days=7 if view=='week' else 30)).isoformat(),next=(start+timedelta(days=7 if view=='week' else 30)).isoformat(),month=start.strftime('%Y-%m'))
@app.route('/shifts/new',methods=['GET','POST'])
@manager_required
def add_shift():return shift_form(None)
@app.route('/shifts/<int:id>/edit',methods=['GET','POST'])
@manager_required
def edit_shift(id):
    with conn() as c:s=c.execute('SELECT * FROM shifts WHERE id=?',(id,)).fetchone()
    if not s:abort(404)
    return shift_form(s)
def shift_form(s):
    if request.method=='POST':
        try:
            data=validate_shift_form(request.form)
            with conn() as c:
                check_shift(c,*data[:4],exclude=s['id'] if s else None)
                if s:c.execute('UPDATE shifts SET employee_id=?,work_date=?,start_time=?,end_time=?,break_minutes=?,station=?,category=?,status=?,notes=?,source="manual",source_key=NULL WHERE id=?',(*data,s['id']))
                else:c.execute('INSERT INTO shifts(employee_id,work_date,start_time,end_time,break_minutes,station,category,status,notes) VALUES(?,?,?,?,?,?,?,?,?)',data)
            flash('Shift saved.','success');return redirect(url_for('schedule'))
        except (ValueError,KeyError,sqlite3.IntegrityError) as e:flash(str(e),'error')
    with conn() as c:employees=c.execute('SELECT * FROM users WHERE active=1 AND role="employee" ORDER BY full_name').fetchall()
    return render_template('shift_form.html',s=s,employees=employees)
@app.post('/shifts/<int:id>/delete')
@manager_required
def delete_shift(id):
    try:
        with conn() as c:c.execute('DELETE FROM shifts WHERE id=?',(id,))
        flash('Shift deleted.','success')
    except sqlite3.IntegrityError:flash('This shift is referenced by a swap request. Resolve the request first.','error')
    return redirect(url_for('schedule'))
@app.route('/employees',methods=['GET','POST'])
@manager_required
def employees():
    if request.method=='POST':
        name=request.form.get('full_name','').strip();username=request.form.get('username','').strip().lower();password=request.form.get('password','');eid=request.form.get('employee_id','').strip() or None
        if not name or not re.fullmatch(r'[a-z0-9_.-]{3,40}',username) or len(password)<10:flash('Enter a name, username (3+ characters), and password (10+ characters).','error')
        else:
            try:
                with conn() as c:c.execute('INSERT INTO users(employee_id,full_name,username,password_hash,role) VALUES(?,?,?,?,?)',(eid,name,username,generate_password_hash(password),'employee'))
                flash('Employee account created.','success')
            except sqlite3.IntegrityError:flash('Username or employee ID already exists.','error')
        return redirect(url_for('employees'))
    with conn() as c:people=c.execute('SELECT id,employee_id,full_name,username,active FROM users WHERE role="employee" ORDER BY active DESC,full_name').fetchall()
    return render_template('employees.html',people=people)
@app.post('/employees/<int:id>/toggle')
@manager_required
def toggle_employee(id):
    with conn() as c:c.execute('UPDATE users SET active=CASE active WHEN 1 THEN 0 ELSE 1 END WHERE id=? AND role="employee"',(id,))
    flash('Employee account updated.','success');return redirect(url_for('employees'))
@app.post('/employees/<int:id>/password')
@manager_required
def employee_password(id):
    pwd=request.form.get('password','')
    if len(pwd)<10:flash('Password must have at least 10 characters.','error')
    else:
        with conn() as c:c.execute('UPDATE users SET password_hash=? WHERE id=? AND role="employee"',(generate_password_hash(pwd),id));flash('Password reset.','success')
    return redirect(url_for('employees'))
@app.route('/time-off',methods=['GET','POST'])
@login_required
def time_off():
    u=user()
    if request.method=='POST':
        if u['role']!='employee':abort(403)
        try:
            a=request.form['start_date'];b=request.form['end_date'];reason=request.form['reason'].strip()
            if parse_date(a)>parse_date(b) or (parse_date(b)-parse_date(a)).days>45 or not reason:raise ValueError('Enter a valid date range (up to 46 days) and reason.')
            with conn() as c:c.execute('INSERT INTO time_off(employee_id,start_date,end_date,reason) VALUES(?,?,?,?)',(u['id'],a,b,reason[:1000]))
            flash('Time-off request submitted.','success')
        except (KeyError,ValueError) as e:flash(str(e),'error')
        return redirect(url_for('time_off'))
    with conn() as c:
        rows=c.execute('''SELECT t.*,u.full_name FROM time_off t JOIN users u ON u.id=t.employee_id '''+('' if u['role']=='manager' else 'WHERE t.employee_id=? ')+'''ORDER BY CASE t.status WHEN 'pending' THEN 0 ELSE 1 END,t.created_at DESC''',() if u['role']=='manager' else (u['id'],)).fetchall()
    return render_template('time_off.html',rows=rows)
@app.post('/time-off/<int:id>/review')
@manager_required
def review_time_off(id):
    decision=request.form.get('decision')
    if decision not in ('approved','declined'):abort(400)
    with conn() as c:
        t=c.execute('SELECT * FROM time_off WHERE id=? AND status="pending"',(id,)).fetchone()
        if not t:flash('Request already reviewed or missing.','error')
        else:
            c.execute('UPDATE time_off SET status=?,manager_comment=?,reviewed_by=?,reviewed_at=CURRENT_TIMESTAMP WHERE id=?',(decision,request.form.get('comment','')[:500],user()['id'],id))
            if decision=='approved':c.execute('UPDATE shifts SET status="Time off" WHERE employee_id=? AND work_date BETWEEN ? AND ?', (t['employee_id'],t['start_date'],t['end_date']))
            flash('Request '+decision+'.','success')
    return redirect(url_for('time_off'))
@app.route('/swaps',methods=['GET','POST'])
@login_required
def swaps():
    u=user()
    if request.method=='POST':
        if u['role']!='employee':abort(403)
        try:
            a=int(request.form['requester_shift_id']);b=int(request.form['target_shift_id'])
            with conn() as c:
                s1=c.execute('SELECT * FROM shifts WHERE id=? AND employee_id=? AND status="Scheduled"',(a,u['id'])).fetchone()
                s2=c.execute('SELECT * FROM shifts WHERE id=? AND employee_id!=? AND status="Scheduled"',(b,u['id'])).fetchone()
                if not s1 or not s2:raise ValueError('Choose your scheduled shift and an eligible coworker shift.')
                if c.execute('SELECT id FROM swaps WHERE status="pending" AND (requester_shift_id IN (?,?) OR target_shift_id IN (?,?))',(a,b,a,b)).fetchone():raise ValueError('One of these shifts already has a pending swap.')
                c.execute('INSERT INTO swaps(requester_id,target_id,requester_shift_id,target_shift_id) VALUES(?,?,?,?)',(u['id'],s2['employee_id'],a,b))
            flash('Swap requested. Your coworker must accept before manager approval.','success')
        except (KeyError,ValueError) as e:flash(str(e),'error')
        return redirect(url_for('swaps'))
    with conn() as c:
        rows=c.execute('''SELECT w.*,a.full_name as requester,b.full_name as target,sa.work_date as a_date,sa.start_time as a_time,sb.work_date as b_date,sb.start_time as b_time FROM swaps w JOIN users a ON a.id=w.requester_id JOIN users b ON b.id=w.target_id JOIN shifts sa ON sa.id=w.requester_shift_id JOIN shifts sb ON sb.id=w.target_shift_id '''+('' if u['role']=='manager' else 'WHERE w.requester_id=? OR w.target_id=? ')+''' ORDER BY CASE w.status WHEN 'pending' THEN 0 ELSE 1 END,w.created_at DESC''',() if u['role']=='manager' else (u['id'],u['id'])).fetchall()
        mine=c.execute('SELECT id,work_date,start_time,station FROM shifts WHERE employee_id=? AND status="Scheduled" AND work_date>=? ORDER BY work_date',(u['id'],date.today().isoformat())).fetchall() if u['role']=='employee' else []
        theirs=c.execute('''SELECT s.id,s.work_date,s.start_time,s.station,u.full_name FROM shifts s JOIN users u ON u.id=s.employee_id WHERE s.employee_id!=? AND u.role="employee" AND u.active=1 AND s.status="Scheduled" AND s.work_date>=? ORDER BY s.work_date''',(u['id'],date.today().isoformat())).fetchall() if u['role']=='employee' else []
    return render_template('swaps.html',rows=rows,mine=mine,theirs=theirs)
@app.post('/swaps/<int:id>/respond')
@login_required
def respond_swap(id):
    decision=request.form.get('decision')
    if decision not in ('accept','decline'):abort(400)
    with conn() as c:
        s=c.execute('SELECT * FROM swaps WHERE id=? AND target_id=? AND status="pending" AND target_accepted IS NULL',(id,user()['id'])).fetchone()
        if not s:abort(403)
        if decision=='decline':c.execute('UPDATE swaps SET status="declined",target_accepted=0 WHERE id=?',(id,))
        else:c.execute('UPDATE swaps SET target_accepted=1 WHERE id=?',(id,))
    flash('Swap response saved.','success');return redirect(url_for('swaps'))
@app.post('/swaps/<int:id>/review')
@manager_required
def review_swap(id):
    decision=request.form.get('decision')
    if decision not in ('approved','declined'):abort(400)
    with conn() as c:
        s=c.execute('SELECT * FROM swaps WHERE id=? AND status="pending"',(id,)).fetchone()
        if not s:flash('Swap already reviewed.','error');return redirect(url_for('swaps'))
        if decision=='approved' and s['target_accepted']!=1:flash('Coworker must accept before approval.','error');return redirect(url_for('swaps'))
        if decision=='approved':
            a=c.execute('SELECT * FROM shifts WHERE id=?',(s['requester_shift_id'],)).fetchone();b=c.execute('SELECT * FROM shifts WHERE id=?',(s['target_shift_id'],)).fetchone()
            if not a or not b or a['employee_id']!=s['requester_id'] or b['employee_id']!=s['target_id'] or a['status']!='Scheduled' or b['status']!='Scheduled':flash('Shifts have changed. Cannot approve.','error');return redirect(url_for('swaps'))
            try:
                check_shift(c,s['target_id'],a['work_date'],a['start_time'],a['end_time'],exclude=b['id'])
                check_shift(c,s['requester_id'],b['work_date'],b['start_time'],b['end_time'],exclude=a['id'])
            except ValueError as e:flash('Swap conflict: '+str(e),'error');return redirect(url_for('swaps'))
            c.execute('UPDATE shifts SET employee_id=?,source="swap",source_key=NULL WHERE id=?',(s['target_id'],a['id']))
            c.execute('UPDATE shifts SET employee_id=?,source="swap",source_key=NULL WHERE id=?',(s['requester_id'],b['id']))
        c.execute('UPDATE swaps SET status=?,manager_comment=?,reviewed_by=?,reviewed_at=CURRENT_TIMESTAMP WHERE id=?',(decision,request.form.get('comment','')[:500],user()['id'],id))
    flash('Swap '+decision+'.','success');return redirect(url_for('swaps'))

# Spreadsheet import is read-only to the SOURCE; user explicitly maps headers, with preview and opt-in import.
def read_spreadsheet(file_bytes,name):
    ext=Path(name.lower()).suffix
    if ext=='.csv':
        text=file_bytes.decode('utf-8-sig');return {'CSV':list(csv.DictReader(io.StringIO(text)))}
    if ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(file_bytes),read_only=True,data_only=True)
        sheets={}
        for ws in wb:
            iterator=ws.iter_rows(values_only=True)
            first=next(iterator, None)
            if first is None:continue
            headers=[str(v).strip() if v is not None else f'Column {i+1}' for i,v in enumerate(first)]
            sheets[ws.title]=[dict(zip(headers, [x.isoformat() if isinstance(x,(date,datetime)) else x for x in row])) for row in iterator if any(v is not None for v in row)]
        wb.close();return sheets
    raise ValueError('Choose a .csv or .xlsx file.')

def parse_import_date(v):
    if isinstance(v,(date,datetime)):return v.date().isoformat() if isinstance(v,datetime) else v.isoformat()
    s=str(v or '').strip()
    for fmt in ('%Y-%m-%d','%m/%d/%Y','%d/%m/%Y','%m/%d/%y','%b %d, %Y','%B %d, %Y'):
        try:return datetime.strptime(s,fmt).date().isoformat()
        except ValueError:pass
    raise ValueError('Cannot read date: '+s)
def parse_import_time(v):
    if isinstance(v,datetime):return v.strftime('%H:%M')
    if hasattr(v,'strftime'):return v.strftime('%H:%M')
    if isinstance(v,(int,float)) and 0<=v<1:return (datetime.min+timedelta(days=v)).strftime('%H:%M')
    s=str(v or '').strip()
    for fmt in ('%H:%M','%H:%M:%S','%I:%M %p','%I %p'):
        try:return datetime.strptime(s,fmt).strftime('%H:%M')
        except ValueError:pass
    raise ValueError('Cannot read time: '+s)
def import_rows(rows,mapping,sheet):
    imported=0;skipped=0;details=[]
    with conn() as c:
        for i,r in enumerate(rows,2):
            try:
                val=lambda key: r.get(mapping.get(key,'none')) if mapping.get(key) not in ('none','') else None
                name=str(val('name') or '').strip();eid=str(val('employee_id') or '').strip() or None
                if not name and not eid:raise ValueError('Missing employee')
                d=parse_import_date(val('date'));st=parse_import_time(val('start'));en=parse_import_time(val('end'))
                rawbreak=val('break');breakmin=int(float(rawbreak or 0));duration(st,en,breakmin)
                employee=c.execute('SELECT * FROM users WHERE employee_id=?',(eid,)).fetchone() if eid else None
                if not employee and name:employee=c.execute('SELECT * FROM users WHERE lower(full_name)=lower(?)',(name,)).fetchone()
                if not employee:raise ValueError('Employee account not found: '+(name or eid)+' (create account first)')
                key=hashlib.sha256(json.dumps([sheet,employee['id'],d,st,en],separators=(',',':')).encode()).hexdigest()
                if c.execute('SELECT id FROM shifts WHERE source_key=?',(key,)).fetchone():skipped+=1;continue
                check_shift(c,employee['id'],d,st,en)
                c.execute('INSERT INTO shifts(employee_id,work_date,start_time,end_time,break_minutes,station,category,notes,source,source_key) VALUES(?,?,?,?,?,?,?,?,?,?)',(employee['id'],d,st,en,breakmin,str(val('station') or '')[:80],str(val('category') or '')[:40],str(val('notes') or '')[:500],'import:'+sheet,key))
                imported+=1
            except Exception as e:
                skipped+=1
                if len(details)<15:details.append(f'Row {i}: {e}')
        c.execute('INSERT INTO imports(filename,sheet_name,imported,skipped,notes) VALUES(?,?,?,?,?)',('Spreadsheet import',sheet,imported,skipped,'; '.join(details)))
    return imported,skipped,details
def load_saved_sheets():
    token=session.get('import_token','')
    if not re.fullmatch(r'[0-9a-f]{32}',token):return {}
    path=BASE/'private_imports'/token
    if not path.is_file():return {}
    try:return read_spreadsheet(path.read_bytes(),session.get('import_filename','sheet.csv'))
    except Exception:return {}

@app.route('/import',methods=['GET','POST'])
@manager_required
def import_page():
    preview=None
    if request.method=='POST':
        action=request.form.get('action')
        if action=='upload':
            f=request.files.get('file')
            if not f or not f.filename:flash('Choose a CSV or XLSX.','error')
            else:
                try:
                    data=f.read();sheets=read_spreadsheet(data,f.filename)
                    # Save source upload locally; do not put sensitive spreadsheet rows in signed browser cookies.
                    upload_dir=BASE/'private_imports';upload_dir.mkdir(exist_ok=True)
                    token=secrets.token_hex(16);(upload_dir/token).write_bytes(data)
                    session['import_token']=token;session['import_filename']=secure_filename(f.filename)
                    if sheets:session['import_sheet']=list(sheets)[0]
                    else:flash('No rows found.','error')
                except Exception as e:flash(str(e),'error')
        elif action=='public_csv':
            try:
                gid=request.form.get('gid','0');gid=str(int(gid))
                req=Request(f'https://docs.google.com/spreadsheets/d/{SHEET_ID}/export?format=csv&gid={gid}',headers={'User-Agent':'Mozilla/5.0'})
                with urlopen(req,timeout=10) as response:blob=response.read(8*1024*1024)
                sheets=read_spreadsheet(blob,'sheet.csv');upload_dir=BASE/'private_imports';upload_dir.mkdir(exist_ok=True)
                token=secrets.token_hex(16);(upload_dir/token).write_bytes(blob);session['import_token']=token;session['import_sheet']='CSV';session['import_filename']='sheet.csv'
            except Exception:flash('Cannot download Google Sheet. Export a CSV/XLSX from Google Sheets and upload it instead.','error')
        elif action=='select':session['import_sheet']=request.form.get('sheet','')
        elif action=='commit':
            sheets=load_saved_sheets();sheet=session.get('import_sheet');mapping={key:request.form.get(key,'none') for key in ('name','employee_id','date','start','end','break','station','category','notes')}
            if any(mapping[x]=='none' for x in ('date','start','end')) or (mapping['name']=='none' and mapping['employee_id']=='none'):flash('Map a name or employee ID, plus date, start, and end.','error')
            elif sheet in sheets:
                imported,skipped,details=import_rows(sheets[sheet],mapping,sheet)
                flash(f'Imported {imported} shifts; skipped {skipped}. '+(' | '.join(details[:3]) if details else ''),'success')
                session.pop('import_token',None);session.pop('import_sheet',None)
                return redirect(url_for('import_page'))
    sheets=load_saved_sheets();selected=session.get('import_sheet');rows=sheets.get(selected,[])
    if selected in sheets:preview={'tabs':list(sheets),'selected':selected,'headers':list(rows[0]) if rows else [],'rows':rows[:5],'count':len(rows)}
    with conn() as c:history=c.execute('SELECT * FROM imports ORDER BY id DESC LIMIT 10').fetchall()
    return render_template('import.html',preview=preview,history=history,sheet_url=f'https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit')
@app.route('/export')
@manager_required
def export_csv():
    rows=shift_rows()
    out=io.StringIO();writer=csv.writer(out);writer.writerow(['Employee ID','Employee','Date','Start','End','Break (min)','Hours','Station','Shift','Status','Notes'])
    for s in rows:writer.writerow([s['employee_code'],s['full_name'],s['work_date'],s['start_time'],s['end_time'],s['break_minutes'],duration(s['start_time'],s['end_time'],s['break_minutes']),s['station'],s['category'],s['status'],s['notes']])
    return app.response_class(out.getvalue(),headers={'Content-Disposition':'attachment; filename=work_schedule.csv'},mimetype='text/csv')
@app.errorhandler(403)
def forbidden(e):return render_template('error.html',message='You do not have permission to access this page.'),403
@app.errorhandler(404)
def not_found(e):return render_template('error.html',message='The page or record was not found.'),404
@app.errorhandler(400)
def bad_request(e):return render_template('error.html',message='Your request was invalid. Refresh the page and try again.'),400
if __name__=='__main__':
    init_db();app.run(host='127.0.0.1',port=int(os.getenv('PORT','5000')),debug=os.getenv('FLASK_DEBUG','0')=='1')
